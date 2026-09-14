"""ShiftGuard — Export builders for Excel (.xlsx), Word (.docx) and PDF."""
import io, os
from datetime import datetime
import config

TEAL, DARK, LIGHT, WHITE = '00A8A8', '1c3a3b', 'E8F8F8', 'FFFFFF'
FONT_DIR  = os.path.join(config.BASE_DIR, 'static', 'fonts')
FONT_REG  = os.path.join(FONT_DIR, 'DejaVuSans.ttf')
FONT_BOLD = os.path.join(FONT_DIR, 'DejaVuSans-Bold.ttf')


def _fname(base, ext):
    return f"{base}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.{ext}"


def _cell_text(v):
    return '' if v is None else str(v)


# ── EXCEL ─────────────────────────────────────────────────────────────────
def excel_table(title, headers, rows, base_filename, col_widths=None, subtitle=None):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (title[:31] or 'Sheet1')
    ncols = max(1, len(headers))
    last_col = get_column_letter(ncols)

    ws.merge_cells(f'A1:{last_col}1')
    c = ws['A1']; c.value = title
    c.font = Font(bold=True, color=WHITE, size=13)
    c.fill = PatternFill('solid', fgColor=DARK)
    c.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 26

    row0 = 2
    if subtitle:
        ws.merge_cells(f'A2:{last_col}2')
        c2 = ws['A2']; c2.value = subtitle
        c2.font = Font(italic=True, size=10, color=DARK)
        c2.alignment = Alignment(horizontal='center')
        row0 = 3

    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=row0, column=ci); cell.value = h
        cell.font = Font(bold=True, color=WHITE, size=10)
        cell.fill = PatternFill('solid', fgColor=TEAL)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for ri, row in enumerate(rows, row0 + 1):
        for ci, val in enumerate(row, 1):
            cell = ws.cell(row=ri, column=ci); cell.value = val
            cell.alignment = Alignment(horizontal='center')
            if ri % 2 == 0: cell.fill = PatternFill('solid', fgColor=LIGHT)

    widths = col_widths or [max(14, len(str(h)) + 4) for h in headers]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf, _fname(base_filename, 'xlsx')


def excel_multi(title, sheets, base_filename):
    """sheets: list of (sheet_title, headers, rows, col_widths|None)."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet_title, headers, rows, col_widths in sheets:
        ws = wb.create_sheet(sheet_title[:31])
        ncols = max(1, len(headers))
        last_col = get_column_letter(ncols)
        ws.merge_cells(f'A1:{last_col}1')
        c = ws['A1']; c.value = title
        c.font = Font(bold=True, color=WHITE, size=12)
        c.fill = PatternFill('solid', fgColor=DARK)
        c.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 24
        for ci, h in enumerate(headers, 1):
            cell = ws.cell(row=2, column=ci); cell.value = h
            cell.font = Font(bold=True, color=WHITE, size=10)
            cell.fill = PatternFill('solid', fgColor=TEAL)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for ri, row in enumerate(rows, 3):
            for ci, val in enumerate(row, 1):
                cell = ws.cell(row=ri, column=ci); cell.value = val
                cell.alignment = Alignment(horizontal='center')
                if ri % 2 == 0: cell.fill = PatternFill('solid', fgColor=LIGHT)
        widths = col_widths or [max(14, len(str(h)) + 4) for h in headers]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf, _fname(base_filename, 'xlsx')


# ── WORD ──────────────────────────────────────────────────────────────────
def word_table(title, headers, rows, base_filename, subtitle=None, sections=None):
    """Single-table Word doc, or pass `sections` = list of (heading, headers, rows)
    to produce a multi-section document instead."""
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

    doc = Document()
    h = doc.add_heading(title, level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in h.runs: run.font.color.rgb = RGBColor(0x1c, 0x3a, 0x3b)
    if subtitle:
        p = doc.add_paragraph(subtitle); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p.runs: p.runs[0].italic = True
    gp = doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    gp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if gp.runs:
        gp.runs[0].font.size = Pt(8); gp.runs[0].font.color.rgb = RGBColor(0x78, 0x90, 0x9c)

    def _add_table(headers, rows):
        table = doc.add_table(rows=1, cols=max(1, len(headers)))
        table.style = 'Light Grid Accent 1'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = table.rows[0].cells
        for i, htext in enumerate(headers):
            hdr[i].text = str(htext)
            for p in hdr[i].paragraphs:
                for r in p.runs: r.font.bold = True
        for row in rows:
            cells = table.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = _cell_text(val)

    if sections:
        for heading, hdrs, rws in sections:
            doc.add_heading(heading, level=2)
            _add_table(hdrs, rws)
            doc.add_paragraph()
    else:
        _add_table(headers, rows)

    buf = io.BytesIO(); doc.save(buf); buf.seek(0)
    return buf, _fname(base_filename, 'docx')


# ── PDF ───────────────────────────────────────────────────────────────────
def _clip(text, col_w):
    max_chars = max(4, int(col_w / 1.7))
    return text if len(text) <= max_chars else text[:max_chars - 1] + '…'


def _pdf_base(orientation='P'):
    from fpdf import FPDF
    pdf = FPDF(orientation=orientation, unit='mm', format='A4')
    pdf.set_auto_page_break(auto=True, margin=12)
    if os.path.exists(FONT_REG) and os.path.exists(FONT_BOLD):
        pdf.add_font('DejaVu', '', FONT_REG)
        pdf.add_font('DejaVu', 'B', FONT_BOLD)
        pdf.set_font('DejaVu', '', 10)
        pdf._sg_font = 'DejaVu'
    else:
        pdf.set_font('Helvetica', '', 10)
        pdf._sg_font = 'Helvetica'
    return pdf


def _pdf_header(pdf, title, subtitle=None):
    pdf.set_font(pdf._sg_font, 'B', 16)
    pdf.set_text_color(28, 58, 59)
    pdf.cell(0, 10, title, new_x='LMARGIN', new_y='NEXT', align='C')
    if subtitle:
        pdf.set_font(pdf._sg_font, '', 10)
        pdf.set_text_color(80, 80, 80)
        pdf.cell(0, 6, subtitle, new_x='LMARGIN', new_y='NEXT', align='C')
    pdf.set_font(pdf._sg_font, '', 8)
    pdf.set_text_color(120, 120, 120)
    pdf.cell(0, 6, f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", new_x='LMARGIN', new_y='NEXT', align='C')
    pdf.ln(2)


def _pdf_table(pdf, headers, rows):
    avail_w = pdf.w - 2 * pdf.l_margin
    col_w = avail_w / max(1, len(headers))

    def draw_head():
        pdf.set_font(pdf._sg_font, 'B', 8)
        pdf.set_fill_color(0, 168, 168); pdf.set_text_color(255, 255, 255)
        for h in headers:
            pdf.cell(col_w, 8, _clip(_cell_text(h), col_w), border=1, align='C', fill=True)
        pdf.ln()

    draw_head()
    pdf.set_font(pdf._sg_font, '', 7.5)
    pdf.set_text_color(30, 30, 30)
    fill = False
    for row in rows:
        if pdf.get_y() > pdf.h - 20:
            pdf.add_page(); draw_head()
            pdf.set_font(pdf._sg_font, '', 7.5); pdf.set_text_color(30, 30, 30)
        pdf.set_fill_color(232, 248, 248) if fill else pdf.set_fill_color(255, 255, 255)
        for val in row:
            pdf.cell(col_w, 7, _clip(_cell_text(val), col_w), border=1, align='C', fill=True)
        pdf.ln(); fill = not fill


def pdf_table(title, headers, rows, base_filename, subtitle=None):
    orientation = 'L' if len(headers) > 6 else 'P'
    pdf = _pdf_base(orientation)
    pdf.add_page()
    _pdf_header(pdf, title, subtitle)
    _pdf_table(pdf, headers, rows)
    buf = io.BytesIO(); buf.write(bytes(pdf.output())); buf.seek(0)
    return buf, _fname(base_filename, 'pdf')


def pdf_multi(title, sections, base_filename, subtitle=None, orientation='L'):
    """sections: list of (heading, headers, rows)."""
    pdf = _pdf_base(orientation)
    pdf.add_page()
    _pdf_header(pdf, title, subtitle)
    for heading, headers, rows in sections:
        pdf.set_font(pdf._sg_font, 'B', 12)
        pdf.set_text_color(0, 168, 168)
        pdf.cell(0, 9, heading, new_x='LMARGIN', new_y='NEXT')
        _pdf_table(pdf, headers, rows)
        pdf.ln(4)
    buf = io.BytesIO(); buf.write(bytes(pdf.output())); buf.seek(0)
    return buf, _fname(base_filename, 'pdf')
