"""Turn an uploaded CV file into content blocks Claude can read.

PDFs and images are sent to Claude as-is, so scanned CVs and complex layouts
work too. Word and text files are converted to plain text locally.
"""
import base64
import os
import re
import shutil
import subprocess

from docx import Document

IMAGE_TYPES = {'png': 'image/png', 'jpg': 'image/jpeg', 'jpeg': 'image/jpeg', 'webp': 'image/webp'}


class ExtractionError(Exception):
    pass


def _b64(path):
    with open(path, 'rb') as f:
        return base64.standard_b64encode(f.read()).decode('ascii')


def _docx_text(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Many CV templates put everything inside tables.
    for table in doc.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                t = cell.text.strip()
                if t and t not in cells:
                    cells.append(t)
            if cells:
                parts.append(' | '.join(cells))
    # Headers often hold the name / contact details.
    for section in doc.sections:
        for p in section.header.paragraphs:
            if p.text.strip():
                parts.insert(0, p.text)
    return '\n'.join(parts)


def _doc_text(path):
    """Legacy .doc — needs antiword or LibreOffice on the server."""
    if shutil.which('antiword'):
        out = subprocess.run(['antiword', path], capture_output=True, text=True, timeout=60)
        if out.returncode == 0 and out.stdout.strip():
            return out.stdout
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if soffice:
        outdir = os.path.dirname(path)
        subprocess.run([soffice, '--headless', '--convert-to', 'docx', '--outdir', outdir, path],
                       capture_output=True, timeout=120)
        converted = os.path.splitext(path)[0] + '.docx'
        if os.path.exists(converted):
            try:
                return _docx_text(converted)
            finally:
                os.remove(converted)
    raise ExtractionError('Old .doc format is not supported on this server — '
                          'please save the file as .docx or PDF and upload again.')


def _rtf_text(path):
    with open(path, 'r', errors='ignore') as f:
        raw = f.read()
    text = re.sub(r'\\par[d]?', '\n', raw)
    text = re.sub(r'\{\\\*[^{}]*\}|\\[a-z]+-?\d* ?|[{}]', '', text)
    return text


def to_content_blocks(path, filename):
    ext = filename.rsplit('.', 1)[-1].lower()

    if ext == 'pdf':
        return [{'type': 'document',
                 'source': {'type': 'base64', 'media_type': 'application/pdf', 'data': _b64(path)}}]

    if ext in IMAGE_TYPES:
        return [{'type': 'image',
                 'source': {'type': 'base64', 'media_type': IMAGE_TYPES[ext], 'data': _b64(path)}}]

    if ext == 'docx':
        text = _docx_text(path)
    elif ext == 'doc':
        text = _doc_text(path)
    elif ext == 'rtf':
        text = _rtf_text(path)
    elif ext == 'txt':
        with open(path, 'r', errors='ignore') as f:
            text = f.read()
    else:
        raise ExtractionError(f'Unsupported file type: .{ext}')

    if not text.strip():
        raise ExtractionError('No text could be read from this file.')
    return [{'type': 'text', 'text': f'<cv filename="{filename}">\n{text}\n</cv>'}]
