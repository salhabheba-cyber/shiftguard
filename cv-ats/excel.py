"""Build the Excel workbook of candidates."""
import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill('solid', fgColor='1F3A5F')
HEADER_FONT = Font(bold=True, color='FFFFFF')
SCORE_FILLS = [
    (80, PatternFill('solid', fgColor='C6EFCE')),  # green
    (65, PatternFill('solid', fgColor='DDEBF7')),  # blue
    (45, PatternFill('solid', fgColor='FFEB9C')),  # yellow
    (0,  PatternFill('solid', fgColor='FFC7CE')),  # red
]

COLUMNS = [
    ('Rank', 7), ('Score', 8), ('Recommendation', 15), ('Full name', 24), ('Email', 28),
    ('Phone', 17), ('Location', 18), ('LinkedIn', 26), ('Current title', 24),
    ('Years exp.', 10), ('Summary', 60), ('Strengths', 45), ('Gaps', 45), ('Skills', 45),
    ('Languages', 18), ('Education', 40), ('Experience', 60), ('Certifications', 30),
    ('CV file', 26), ('Uploaded', 18),
]


def _lines(items):
    return '\n'.join(f'• {i}' for i in items)


def _sheet_title(text, used):
    title = re.sub(r'[\[\]\*\?/\\:]', '', text)[:28] or 'Job'
    base, n = title, 2
    while title in used:
        title = f'{base[:25]}_{n}'
        n += 1
    used.add(title)
    return title


def _fill_job_sheet(ws, job, candidates):
    ws.append([c for c, _ in COLUMNS])
    for i, (_, width) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
        cell = ws.cell(row=1, column=i)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.freeze_panes = 'E2'

    done = [c for c in candidates if c['status'] == 'done']
    for rank, c in enumerate(done, 1):
        d = c['data']
        ws.append([
            rank, c['score'], c['recommendation'], d.get('full_name'), d.get('email'),
            d.get('phone'), d.get('location'), d.get('linkedin'), d.get('current_title'),
            d.get('total_years_experience'), d.get('summary'),
            _lines(d.get('strengths', [])), _lines(d.get('gaps', [])),
            ', '.join(d.get('skills', [])), ', '.join(d.get('languages', [])),
            _lines(f"{e['degree']} — {e['institution']} ({e['year']})" for e in d.get('education', [])),
            _lines(f"{e['title']} @ {e['company']} ({e['period']})" for e in d.get('experience', [])),
            ', '.join(d.get('certifications', [])), c['filename'], c['created_at'],
        ])
        row = ws.max_row
        for cell in ws[row]:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        for threshold, fill in SCORE_FILLS:
            if (c['score'] or 0) >= threshold:
                ws.cell(row=row, column=2).fill = fill
                ws.cell(row=row, column=3).fill = fill
                break

    if len(done) >= 1:
        ws.auto_filter.ref = f'A1:{get_column_letter(len(COLUMNS))}{ws.max_row}'

    pending = [c for c in candidates if c['status'] != 'done']
    if pending:
        ws.append([])
        ws.append(['Not analysed', '', '', 'File', 'Status / error'])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        for c in pending:
            ws.append(['', '', '', c['filename'], c['error'] or c['status']])


def _fill_criteria_sheet(ws, candidates):
    ws.append(['Full name', 'Overall score', 'Criterion', 'Score /10', 'Comment'])
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for col, width in zip('ABCDE', (24, 12, 40, 10, 70)):
        ws.column_dimensions[col].width = width
    for c in candidates:
        if c['status'] != 'done':
            continue
        for cs in c['data'].get('criteria_scores', []):
            ws.append([c['name'], c['score'], cs['criterion'], cs['score'], cs['comment']])
            for cell in ws[ws.max_row]:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
    ws.freeze_panes = 'A2'


def build_workbook(jobs_with_candidates):
    """jobs_with_candidates: list of (job_row, [candidate dicts])."""
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for job, candidates in jobs_with_candidates:
        _fill_job_sheet(wb.create_sheet(_sheet_title(job['title'], used)), job, candidates)
        if len(jobs_with_candidates) == 1:
            _fill_criteria_sheet(wb.create_sheet('Criteria scores'), candidates)
    if not wb.sheetnames:
        wb.create_sheet('Candidates').append(['No jobs yet'])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
